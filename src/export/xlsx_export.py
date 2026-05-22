"""Convert table JSON to Excel workbook (one sheet per table)."""

from __future__ import annotations

import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from src.models.schema import DocumentTables, TableRecord


def _safe_sheet_name(table: TableRecord) -> str:
    """Excel sheet names max 31 chars, no : \\ / ? * [ ]"""
    base = table.title.strip() or f"Table_{table.table_id}"
    base = re.sub(r"[\:\\/?*\[\]]", "_", base)[:28]
    return f"{base}_{table.table_id}"[:31]


def write_tables_xlsx(doc: DocumentTables, output_dir: Path) -> Path:
    """Write <document_id>_tables.xlsx with one worksheet per table."""
    output_dir.mkdir(parents=True, exist_ok=True)
    out_path = output_dir / f"{doc.document_id}_tables.xlsx"
    print(f"[export/xlsx] Writing workbook: {out_path}")

    wb = Workbook()
    # Remove default sheet; we add per table
    default = wb.active
    wb.remove(default)

    if not doc.tables:
        ws = wb.create_sheet("NoTables")
        ws.append(["message", "No tables extracted"])
        print("[export/xlsx] No tables — wrote placeholder sheet")
    else:
        for table in doc.tables:
            sheet_name = _safe_sheet_name(table)
            print(f"[export/xlsx]   Sheet: {sheet_name} ({len(table.rows)} rows)")
            ws = wb.create_sheet(sheet_name)
            ws.append(table.columns)
            for row in table.rows:
                # Pad or trim row to match column count
                cells = list(row)
                ncol = len(table.columns)
                if len(cells) < ncol:
                    cells.extend([""] * (ncol - len(cells)))
                elif len(cells) > ncol:
                    cells = cells[:ncol]
                ws.append(cells)
            _autosize_columns(ws, len(table.columns))

    wb.save(out_path)
    print(f"[export/xlsx] Saved workbook with {max(len(doc.tables), 1)} sheet(s)")
    return out_path


def _autosize_columns(ws, num_cols: int) -> None:
    """Rough column width from content length."""
    for col_idx in range(1, num_cols + 1):
        letter = get_column_letter(col_idx)
        max_len = 10
        for cell in ws[letter]:
            if cell.value is not None:
                max_len = max(max_len, min(len(str(cell.value)), 50))
        ws.column_dimensions[letter].width = max_len + 2
